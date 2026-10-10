"""日別客数予測ツール (Excel) を生成するスクリプト。

使い方:
    python build_forecast_tool.py 出力.xlsx [先月データ.xlsx] [先々月データ.xlsx] [前年当月データ.xlsx]

データファイルを渡すと、その内容を貼り付け済みの状態で出力する（動作確認用）。
渡さなければ空のテンプレートになる。
"""
import sys

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Meiryo UI"
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FT = Font(name=FONT, size=14, bold=True)
FIN = Font(name=FONT, size=10, color="0000FF")
FW = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FG = Font(name=FONT, size=9, color="808080")
FNOTE = Font(name=FONT, size=9, color="595959")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_HD = PatternFill("solid", fgColor="44546A")
FILL_SUB = PatternFill("solid", fgColor="D9E1F2")
FILL_GR = PatternFill("solid", fgColor="EDEDED")
FILL_KPI = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")

TYPES = ["月", "火", "水", "木", "金", "土", "日", "祝"]
HOURS = list(range(11, 24))
DATA_ROWS = 800  # 31日×24時間=744行 + 余裕
DS1, DS2, DSP = "先月データ", "先々月データ", "前年当月データ"
LD, DO, CL_ = "ランチ+ディナー", "ディナーのみ", "休業"
DSO = "売上予測貼付用"
BANDS = ["ランチ", "アイドル", "ディナー", "ナイト"]  # 【事前】売上予測 の列順
# 時刻→時間帯（基幹システムの時間帯別売上実績表の区分）
BAND_OF_HOUR = {h: ("ナイト" if h <= 7 or h >= 22 else "モーニング" if h <= 10 else "ランチ" if h <= 14
                    else "アイドル" if h <= 17 else "ディナー") for h in range(24)}

HOLIDAYS = [
    ("2025-01-01", "元日"), ("2025-01-13", "成人の日"), ("2025-02-11", "建国記念の日"),
    ("2025-02-23", "天皇誕生日"), ("2025-02-24", "振替休日"), ("2025-03-20", "春分の日"),
    ("2025-04-29", "昭和の日"), ("2025-05-03", "憲法記念日"), ("2025-05-04", "みどりの日"),
    ("2025-05-05", "こどもの日"), ("2025-05-06", "振替休日"), ("2025-07-21", "海の日"),
    ("2025-08-11", "山の日"), ("2025-09-15", "敬老の日"), ("2025-09-23", "秋分の日"),
    ("2025-10-13", "スポーツの日"), ("2025-11-03", "文化の日"), ("2025-11-23", "勤労感謝の日"),
    ("2025-11-24", "振替休日"),
    ("2026-01-01", "元日"), ("2026-01-12", "成人の日"), ("2026-02-11", "建国記念の日"),
    ("2026-02-23", "天皇誕生日"), ("2026-03-20", "春分の日"), ("2026-04-29", "昭和の日"),
    ("2026-05-03", "憲法記念日"), ("2026-05-04", "みどりの日"), ("2026-05-05", "こどもの日"),
    ("2026-05-06", "振替休日"), ("2026-07-20", "海の日"), ("2026-08-11", "山の日"),
    ("2026-09-21", "敬老の日"), ("2026-09-22", "国民の休日"), ("2026-09-23", "秋分の日"),
    ("2026-10-12", "スポーツの日"), ("2026-11-03", "文化の日"), ("2026-11-23", "勤労感謝の日"),
    ("2027-01-01", "元日"), ("2027-01-11", "成人の日"), ("2027-02-11", "建国記念の日"),
    ("2027-02-23", "天皇誕生日"), ("2027-03-21", "春分の日"), ("2027-03-22", "振替休日"),
    ("2027-04-29", "昭和の日"), ("2027-05-03", "憲法記念日"), ("2027-05-04", "みどりの日"),
    ("2027-05-05", "こどもの日"), ("2027-07-19", "海の日"), ("2027-08-11", "山の日"),
    ("2027-09-20", "敬老の日"), ("2027-09-23", "秋分の日"), ("2027-10-11", "スポーツの日"),
    ("2027-11-03", "文化の日"), ("2027-11-23", "勤労感謝の日"),
]


def wd_formula(cell):
    return f'CHOOSE(WEEKDAY({cell},2),"月","火","水","木","金","土","日")'


def type_formula(cell):
    """日付セルから日タイプ（月〜日 / 祝）を返す式。"""
    return f'IF(COUNTIF(祝日!$A:$A,{cell})>0,"祝",{wd_formula(cell)})'


def style(c, font=F, fill=None, fmt=None, align=None, border=True):
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if border:
        c.border = BOX


def header(ws, row, col, labels, fill=FILL_HD, font=FW):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        style(c, font=font, fill=fill, align=Alignment(horizontal="center", vertical="center", wrap_text=True))


# ---- 集計シートのレイアウト ----
D0, D1 = 3, 33  # 日別行
# 先月/先々月ブロック内の列オフセット
BLK = ["日付", "タイプ", "曜日", "客数", "売上", "前年客数", "ランチ客数", "ディナー客数",
       "ランチ営業", "前年ディナー", "前年日=祝", "前年売上"]
B1, B2 = 2, 14  # 先月ブロック開始列(B)、先々月ブロック開始列(N)
BP = 26  # 前年当月ブロック開始列(Z)
BLKP = ["前年当月 日付", "タイプ", "客数", "ランチ客数", "ディナー客数", "売上"]


def bcol(start, name):
    return CL(start + BLK.index(name))


def pcol(name):
    return CL(BP + BLKP.index(name))


MAPH, MAPB = "設定!$E$4:$E$27", "設定!$F$4:$F$27"


def band_formula(cell):
    return f'IFERROR(INDEX({MAPB},MATCH({cell},{MAPH},0)),"")'


def rng(col):
    return f"集計!${col}${D0}:${col}${D1}"


def build(out, src1=None, src2=None, srcp=None):
    wb = openpyxl.Workbook()
    ws_help = wb.active
    ws_help.title = "使い方"
    ws_fc = wb.create_sheet("予測")
    ws_hr = wb.create_sheet("時間帯別予測")
    ws_set = wb.create_sheet("設定")
    ws_d1 = wb.create_sheet(DS1)
    ws_d2 = wb.create_sheet(DS2)
    ws_dp = wb.create_sheet(DSP)
    ws_out = wb.create_sheet(DSO, 3)
    ws_hol = wb.create_sheet("祝日")
    ws_c = wb.create_sheet("集計")

    # ---------------- 設定（先にセル位置を決める） ----------------
    S = {}
    s = ws_set
    s["A1"] = "設定"
    s["A1"].font = FT
    s.column_dimensions["A"].width = 32
    s.column_dimensions["B"].width = 16
    s.column_dimensions["C"].width = 80
    m1 = f"'{DS1}'!A2"
    m2 = f"'{DS2}'!A2"
    mp = f"'{DSP}'!A2"
    rows = [
        ("sec", "▼ 対象月"),
        ("m1", "先月（データから自動判定）", f'=IF(ISNUMBER({m1}),DATE(YEAR({m1}),MONTH({m1}),1),"")', "yyyy年m月", False, "先月データの1行目の日付から自動判定"),
        ("m2", "先々月（データから自動判定）", f'=IF(ISNUMBER({m2}),DATE(YEAR({m2}),MONTH({m2}),1),"")', "yyyy年m月", False, "先々月データの1行目の日付から自動判定"),
        ("mp", "前年当月（データから自動判定）", f'=IF(ISNUMBER({mp}),DATE(YEAR({mp}),MONTH({mp}),1),"")', "yyyy年m月", False, "前年当月データの1行目の日付から自動判定（例：当月が2026年10月なら2025年10月を指定して出力）"),
        ("cur", "予測する月（当月）", None, "yyyy年m月", True, "通常は先月の翌月（自動）。別の月にしたい場合は 2026/11/1 のように月初日を直接入力"),
        ("days", "当月の日数", None, "0", False, ""),
        ("sec", "▼ 営業時間"),
        ("dinner", "夜営業の開始時刻", 17, '0"時"', True, "この時刻以降をディナー、より前をランチ（昼〜夕方の通し営業を含む）として分けて計算"),
        ("thr", "ランチ営業ありと判定する客数", 10, '0"人"', True, "過去データで、夜営業開始前の客数がこれ以上の日を「ランチ営業した日」とみなす"),
        ("sec", "▼ 重み"),
        ("w1", "先月の重み", 0.6, "0.00", True, "直近ほど重視。先々月と合計1でなくてもOK（比率で按分）"),
        ("w2", "先々月の重み", 0.4, "0.00", True, ""),
        ("wp", "前年同曜日の重み（0〜1）", 0.5, "0.00", True, "前年当月データがある日は「直近ベース×(1−重み)＋前年同曜日×直近前年比×重み」。0にすると前年の日別値は使わない"),
        ("sec", "▼ 補正"),
        ("trend", "トレンド係数", 1.0, "0.00", True, "全体の上げ下げ。例：販促・値上げで5%増を見込む→1.05"),
        ("smanual", "季節係数（前年当月データが無い場合の手入力）", 1.0, "0.000", True, "前年当月データを貼ればそちらから自動計算されるので入力不要"),
        ("season", "季節係数（適用値）", None, "0.000", False, "＝前年当月の曜日別ディナー平均 ÷ 前年の先月・先々月の曜日別ディナー平均（祝日を除く）"),
        ("yoy", "直近の前年比（適用値）", None, "0.000", False, "先月の客数 ÷ 先月の前年客数（先月が無ければ先々月）。前年同曜日の値に掛けて今年の水準に合わせる"),
        ("sec", "▼ 客単価の補正"),
        ("ptrend", "客単価トレンド係数", 1.0, "0.00", True, "値上げ・メニュー改定・クーポン施策など。例：3%値上げ→1.03"),
        ("psmanual", "客単価の季節係数（前年当月データが無い場合の手入力）", 1.0, "0.000", True, "前年当月データを貼ればそちらから自動計算"),
        ("pseason", "客単価の季節係数（適用値）", None, "0.000", False, "＝前年当月の客単価 ÷ 前年の先月・先々月の客単価（重み付け）"),
        ("pmin", "時間帯別客単価を計算する最低客数", 10, '0"人"', True, "この人数未満の時間帯は、単価がぶれるので全体（全曜日合算）の単価で代用"),
        ("sec", "▼ 判断材料"),
        ("yoy1", "【参考】先月 客数前年比", None, "0.0%", False, ""),
        ("yoy2", "【参考】先々月 客数前年比", None, "0.0%", False, ""),
        ("pc1", "【参考】先月 客単価", None, "#,##0", False, ""),
        ("pc2", "【参考】先々月 客単価", None, "#,##0", False, ""),
        ("ppy", "【参考】前年当月 客単価", None, "#,##0", False, ""),
        ("ppy1", "【参考】先月の前年 客単価", None, "#,##0", False, "データの「売上前年」÷「客数前年」"),
        ("ppy2", "【参考】先々月の前年 客単価", None, "#,##0", False, ""),
        ("check", "データチェック", None, None, False, ""),
    ]
    r = 3
    for row in rows:
        if row[0] == "sec":
            s.cell(row=r, column=1, value=row[1]).font = FB
            r += 1
            continue
        key, label, val, fmt, is_in, note = row
        S[key] = f"設定!$B${r}"
        style(s.cell(row=r, column=1, value=label), font=FB, fill=FILL_SUB)
        style(s.cell(row=r, column=2, value=val), font=FIN if is_in else F, fill=FILL_IN if is_in else None, fmt=fmt, align=CENTER)
        s.cell(row=r, column=3, value=note).font = FNOTE
        r += 1
    s.cell(row=r + 1, column=1, value="黄色セル＝入力欄（青字）。その他は自動計算です。").font = FNOTE

    s.column_dimensions["E"].width = 8
    s.column_dimensions["F"].width = 12
    s["E2"] = "▼ 時刻→時間帯の対応表（【事前】売上予測の列分けに使用・変更可）"
    s["E2"].font = FB
    header(s, 3, 5, ["時刻", "時間帯"])
    for h_ in range(24):
        style(s.cell(row=4 + h_, column=5, value=h_), fmt='0"時"', align=CENTER)
        style(s.cell(row=4 + h_, column=6, value=BAND_OF_HOUR[h_]), font=FIN, fill=FILL_IN, align=CENTER)
    bdv = DataValidation(type="list", formula1='"' + ",".join(["モーニング"] + BANDS) + '"', allow_blank=True)
    s.add_data_validation(bdv)
    bdv.add("F4:F27")
    s.column_dimensions["C"].width = 70

    def setcell(key, formula):
        s[S[key].split("!")[1].replace("$", "")] = formula

    setcell("cur", f'=IF({S["m1"]}="","",EDATE({S["m1"]},1))')
    setcell("days", f'=IF({S["cur"]}="","",DAY(EOMONTH({S["cur"]},0)))')

    # ---------------- データ貼り付けシート ----------------
    for ws, src, label in ((ws_d1, src1, "先月"), (ws_d2, src2, "先々月"), (ws_dp, srcp, "前年当月（任意）")):
        ws.sheet_properties.tabColor = "FFC000"
        if src:
            sws = openpyxl.load_workbook(src).active
            for row in sws.iter_rows(values_only=True):
                ws.append(list(row))
            for row in ws.iter_rows(min_row=2, max_col=2):
                for c in row:
                    c.number_format = "yyyy/mm/dd"
        else:
            ws["A1"] = f"← A1セルを選択して、{label}の「時間帯別売上実績表」を見出し行ごと貼り付けてください"
            ws["A1"].font = Font(name=FONT, size=11, bold=True, color="C00000")
        ws["Z1"] = "日タイプ(自動)"
        style(ws["Z1"], font=FB, fill=FILL_GR)
        for rr in range(2, DATA_ROWS + 1):
            ws.cell(row=rr, column=26, value=f'=IF(ISNUMBER(A{rr}),{type_formula(f"A{rr}")},"")').font = FG
        ws.column_dimensions["Z"].width = 14
        if ws is not ws_dp:
            blk = B1 if ws is ws_d1 else B2
            ws["AA1"] = "ランチ営業日(自動)"
            style(ws["AA1"], font=FB, fill=FILL_GR)
            dcol, fcol = bcol(blk, "日付"), bcol(blk, "ランチ営業")
            for rr in range(2, DATA_ROWS + 1):
                ws.cell(row=rr, column=27, value=(
                    f'=IF(ISNUMBER(A{rr}),IFERROR(INDEX({rng(fcol)},MATCH(A{rr},{rng(dcol)},0)),0),"")')).font = FG
            ws.column_dimensions["AA"].width = 16
            ws["AB1"] = "時間帯区分(自動)"
            style(ws["AB1"], font=FB, fill=FILL_GR)
            for rr in range(2, DATA_ROWS + 1):
                ws.cell(row=rr, column=28, value=f'=IF(ISNUMBER(A{rr}),{band_formula(f"K{rr}")},"")').font = FG
            ws.column_dimensions["AB"].width = 16
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 12
        ws.freeze_panes = "A2"
    ws_dp["AC1"] = "例：当月が2026年10月なら、期間を2025/10/1〜2025/11/1で出力したものを貼り付け（「当年」列を前年データとして使います。月末の前年同曜日は翌月1日になるため1日多めに）"
    ws_dp["AC1"].font = FB

    # ---------------- 祝日 ----------------
    header(ws_hol, 1, 1, ["日付", "名称"])
    for i, (d, name) in enumerate(HOLIDAYS, start=2):
        y, m, dd = map(int, d.split("-"))
        ws_hol.cell(row=i, column=1, value=f"=DATE({y},{m},{dd})")
        ws_hol.cell(row=i, column=2, value=name)
        style(ws_hol.cell(row=i, column=1), fmt="yyyy/mm/dd(aaa)")
        style(ws_hol.cell(row=i, column=2))
    ws_hol["D1"] = "祝日扱いにしたい日（お盆・年末年始・地域イベント等）は、この表の下に日付を追加してください。前年分の祝日も前年データの判定に使います。"
    ws_hol["D1"].font = FB
    ws_hol.column_dimensions["A"].width = 16
    ws_hol.column_dimensions["B"].width = 16

    # ---------------- 集計：日別 ----------------
    c = ws_c
    c.sheet_properties.tabColor = "A6A6A6"
    c["A1"] = "集計（自動計算・編集不要）"
    c["A1"].font = FT
    header(c, 2, 1, ["日"])
    header(c, 2, B1, ["先月 " + BLK[0]] + BLK[1:])
    header(c, 2, B2, ["先々月 " + BLK[0]] + BLK[1:])
    header(c, 2, BP, BLKP)
    dn = S["dinner"]
    for i in range(31):
        r = D0 + i
        c.cell(row=r, column=1, value=i + 1)
        style(c.cell(row=r, column=1), align=CENTER)
        for blk, sheet, mkey in ((B1, DS1, "m1"), (B2, DS2, "m2")):
            mc = S[mkey]
            d = f"{bcol(blk, '日付')}{r}"
            q = f"'{sheet}'!$Q:$Q"
            rr_ = f"'{sheet}'!$R:$R"
            a = f"'{sheet}'!$A:$A"
            k = f"'{sheet}'!$K:$K"
            lunch_crit = f'{k},">=6",{k},"<"&{dn}'
            vals = {
                "日付": f'=IF({mc}="","",IF($A{r}<=DAY(EOMONTH({mc},0)),{mc}+$A{r}-1,""))',
                "タイプ": f'=IF({d}="","",{type_formula(d)})',
                "曜日": f'=IF({d}="","",{wd_formula(d)})',
                "客数": f'=IF({d}="","",SUMIFS({q},{a},{d}))',
                "売上": f"=IF({d}=\"\",\"\",SUMIFS('{sheet}'!$L:$L,{a},{d}))",
                "前年客数": f'=IF({d}="","",SUMIFS({rr_},{a},{d}))',
                "ランチ客数": f'=IF({d}="","",SUMIFS({q},{a},{d},{lunch_crit}))',
                "ディナー客数": f'=IF({d}="","",{bcol(blk, "客数")}{r}-{bcol(blk, "ランチ客数")}{r})',
                "ランチ営業": f'=IF({d}="","",IF({bcol(blk, "ランチ客数")}{r}>={S["thr"]},1,0))',
                "前年ディナー": f'=IF({d}="","",{bcol(blk, "前年客数")}{r}-SUMIFS({rr_},{a},{d},{lunch_crit}))',
                "前年日=祝": f'=IF({d}="","",IF(COUNTIF(祝日!$A:$A,{d}-364)>0,1,0))',
                "前年売上": f"=IF({d}=\"\",\"\",SUMIFS('{sheet}'!$M:$M,{a},{d}))",
            }
            for j, name in enumerate(BLK):
                cc = c.cell(row=r, column=blk + j, value=vals[name])
                style(cc, fmt="m/d(aaa)" if name == "日付" else "#,##0", align=CENTER if j < 3 or name in ("ランチ営業", "前年日=祝") else None)
        # 前年当月
        mp_ = S["mp"]
        d = f"{pcol('前年当月 日付')}{r}"
        q, a, k = f"'{DSP}'!$Q:$Q", f"'{DSP}'!$A:$A", f"'{DSP}'!$K:$K"
        pv = {
            "前年当月 日付": f'=IF({mp_}="","",IF($A{r}<=DAY(EOMONTH({mp_},0)),{mp_}+$A{r}-1,""))',
            "タイプ": f'=IF({d}="","",{type_formula(d)})',
            "客数": f'=IF({d}="","",SUMIFS({q},{a},{d}))',
            "ランチ客数": f'=IF({d}="","",SUMIFS({q},{a},{d},{k},">=6",{k},"<"&{dn}))',
            "ディナー客数": f'=IF({d}="","",{pcol("客数")}{r}-{pcol("ランチ客数")}{r})',
            "売上": f"=IF({d}=\"\",\"\",SUMIFS('{DSP}'!$L:$L,{a},{d}))",
        }
        for j, name in enumerate(BLKP):
            style(c.cell(row=r, column=BP + j, value=pv[name]), fmt="m/d(aaa)" if j == 0 else "#,##0",
                  align=CENTER if j < 2 else None)

    # ---------------- 集計：日タイプ別 ----------------
    TR = 38
    c.cell(row=TR - 2, column=1, value="日タイプ別 基準値（ディナー＝営業日平均、ランチ＝ランチ営業した日の平均）").font = FB
    thdr = ["タイプ", "先月 営業日数", "先月 ランチ営業日数", "先月 ディナー平均", "先月 ランチ平均",
            "先々月 営業日数", "先々月 ランチ営業日数", "先々月 ディナー平均", "先々月 ランチ平均",
            "重み付け ディナー", "重み付け ランチ", "採用 ディナー基準", "採用 ランチ基準",
            "先月 客単価", "先々月 客単価", "採用 客単価",
            "前年先月 ディナー平均", "前年先々月 ディナー平均", "前年(直近) 重み付け", "前年当月 ディナー平均", "季節係数に使用"]
    header(c, TR - 1, 1, thdr)
    w1, w2 = S["w1"], S["w2"]
    POOL = TR + 9  # 補助行
    sun = TR + TYPES.index("日")

    def wavg(x1, n1, x2, n2):
        return (f'IF(AND({n1}>0,{n2}>0),({w1}*{x1}+{w2}*{x2})/({w1}+{w2}),'
                f'IF({n1}>0,{x1},IF({n2}>0,{x2},"")))')

    for i, t in enumerate(TYPES):
        r = TR + i
        A = f"$A{r}"
        cells = {}
        for blk, off in ((B1, 0), (B2, 4)):
            ty, cu, ln, dn_, fl = (rng(bcol(blk, n)) for n in ("タイプ", "客数", "ランチ客数", "ディナー客数", "ランチ営業"))
            cells[2 + off] = f'=COUNTIFS({ty},{A},{cu},">0")'
            cells[3 + off] = f'=COUNTIFS({ty},{A},{fl},1)'
            n_col, nl_col = CL(2 + off), CL(3 + off)
            cells[4 + off] = f'=IF({n_col}{r}>0,AVERAGEIFS({dn_},{ty},{A},{cu},">0"),"")'
            cells[5 + off] = f'=IF({nl_col}{r}>0,AVERAGEIFS({ln},{ty},{A},{fl},1),"")'
        cells[10] = "=" + wavg(f"D{r}", f"B{r}", f"H{r}", f"F{r}")
        cells[11] = "=" + wavg(f"E{r}", f"C{r}", f"I{r}", f"G{r}")
        cells[12] = f'=IF(J{r}<>"",J{r},IF(AND({A}="祝",$J${sun}<>""),$J${sun},$L${POOL}))'
        if t in ("月", "火", "水", "木", "金"):
            cells[13] = f'=IF(K{r}<>"",K{r},IF($M${POOL + 1}<>"",$M${POOL + 1},N($M${POOL})))'
        else:
            cells[13] = f'=IF(K{r}<>"",K{r},IF(AND({A}="祝",$K${sun}<>""),$K${sun},N($M${POOL})))'
        for blk, col in ((B1, 14), (B2, 15)):
            ty, cu, sa = (rng(bcol(blk, n)) for n in ("タイプ", "客数", "売上"))
            cells[col] = f'=IFERROR(SUMIFS({sa},{ty},{A})/SUMIFS({cu},{ty},{A}),"")'
        cells[16] = f'=IF(AND(N{r}<>"",O{r}<>""),({w1}*N{r}+{w2}*O{r})/({w1}+{w2}),IF(N{r}<>"",N{r},IF(O{r}<>"",O{r},$P${POOL})))'
        if t != "祝":
            for blk, col in ((B1, 17), (B2, 18)):
                wd_, pd_, ph = (rng(bcol(blk, n)) for n in ("曜日", "前年ディナー", "前年日=祝"))
                cells[col] = f'=IFERROR(AVERAGEIFS({pd_},{wd_},{A},{pd_},">0",{ph},0),"")'
            cells[19] = (f'=IF(AND(Q{r}<>"",R{r}<>""),({w1}*Q{r}+{w2}*R{r})/({w1}+{w2}),'
                         f'IF(Q{r}<>"",Q{r},IF(R{r}<>"",R{r},"")))')
            cells[20] = f'=IFERROR(AVERAGEIFS({rng(pcol("ディナー客数"))},{rng(pcol("タイプ"))},{A},{rng(pcol("ディナー客数"))},">0"),"")'
            cells[21] = f'=IF(AND(S{r}<>"",T{r}<>""),1,0)'
        c.cell(row=r, column=1, value=t)
        for col, v in cells.items():
            c.cell(row=r, column=col, value=v)
        for col in range(1, 22):
            fmt = "0" if col in (2, 3, 6, 7, 21) else "#,##0" if col in (14, 15, 16) else "#,##0.0"
            style(c.cell(row=r, column=col), fmt=fmt, align=CENTER if col == 1 else None)
    # 補助行（データが無いタイプの代用値）
    d1c, d2c = rng(bcol(B1, "客数")), rng(bcol(B2, "客数"))
    pool_rows = [
        (POOL, "全体平均",
         {12: "=IFERROR(" + wavg(f'AVERAGEIF({rng(bcol(B1, "客数"))},">0")', f'COUNTIF({d1c},">0")',
                                 f'AVERAGEIF({rng(bcol(B2, "客数"))},">0")', f'COUNTIF({d2c},">0")') + ",0)",
          13: "=IFERROR(" + wavg(f'AVERAGEIF({rng(bcol(B1, "ランチ営業"))},1,{rng(bcol(B1, "ランチ客数"))})',
                                 f'COUNTIF({rng(bcol(B1, "ランチ営業"))},1)',
                                 f'AVERAGEIF({rng(bcol(B2, "ランチ営業"))},1,{rng(bcol(B2, "ランチ客数"))})',
                                 f'COUNTIF({rng(bcol(B2, "ランチ営業"))},1)') + ',"")',
          16: f'=IFERROR((SUM({rng(bcol(B1, "売上"))})+SUM({rng(bcol(B2, "売上"))}))/(SUM({d1c})+SUM({d2c})),0)'}),
        (POOL + 1, "平日ランチ平均", {}),
    ]
    # 平日（月〜金・祝日除く）のランチ営業日平均
    parts = []
    for blk in (B1, B2):
        ty, ln, fl = (rng(bcol(blk, n)) for n in ("タイプ", "ランチ客数", "ランチ営業"))
        tot = f'(SUMIFS({ln},{fl},1)-SUMIFS({ln},{fl},1,{ty},"土")-SUMIFS({ln},{fl},1,{ty},"日")-SUMIFS({ln},{fl},1,{ty},"祝"))'
        cnt = f'(COUNTIFS({fl},1)-COUNTIFS({fl},1,{ty},"土")-COUNTIFS({fl},1,{ty},"日")-COUNTIFS({fl},1,{ty},"祝"))'
        parts.append((tot, cnt))
    pool_rows[1][2][13] = "=IFERROR(" + wavg(f"{parts[0][0]}/{parts[0][1]}", parts[0][1], f"{parts[1][0]}/{parts[1][1]}", parts[1][1]) + ',"")'
    for r, label, cells in pool_rows:
        c.cell(row=r, column=1, value=label)
        for col, v in cells.items():
            c.cell(row=r, column=col, value=v)
        for col in range(1, 22):
            style(c.cell(row=r, column=col), font=FB, fill=FILL_GR, fmt="#,##0" if col == 16 else "#,##0.0")

    # 季節係数・直近前年比
    setcell("season", f'=IF(AND({S["mp"]}<>"",SUM(集計!$U${TR}:$U${TR + 6})>0),'
                      f'SUMIFS(集計!$T${TR}:$T${TR + 6},集計!$U${TR}:$U${TR + 6},1)/SUMIFS(集計!$S${TR}:$S${TR + 6},集計!$U${TR}:$U${TR + 6},1),'
                      f'{S["smanual"]})')
    setcell("yoy1", f'=IFERROR(SUM({rng(bcol(B1, "客数"))})/SUM({rng(bcol(B1, "前年客数"))}),"")')
    setcell("yoy2", f'=IFERROR(SUM({rng(bcol(B2, "客数"))})/SUM({rng(bcol(B2, "前年客数"))}),"")')
    setcell("yoy", f'=IF({S["yoy1"]}<>"",{S["yoy1"]},IF({S["yoy2"]}<>"",{S["yoy2"]},1))')
    def ratio(num, den):
        return f'IFERROR(SUM({rng(num)})/SUM({rng(den)}),"")'

    setcell("pc1", "=" + ratio(bcol(B1, "売上"), bcol(B1, "客数")))
    setcell("pc2", "=" + ratio(bcol(B2, "売上"), bcol(B2, "客数")))
    setcell("ppy", "=" + ratio(pcol("売上"), pcol("客数")))
    setcell("ppy1", "=" + ratio(bcol(B1, "前年売上"), bcol(B1, "前年客数")))
    setcell("ppy2", "=" + ratio(bcol(B2, "前年売上"), bcol(B2, "前年客数")))
    p1, p2 = S["ppy1"], S["ppy2"]
    setcell("pseason", (f'=IFERROR(IF({S["ppy"]}="",{S["psmanual"]},{S["ppy"]}/'
                        f'IF(AND(ISNUMBER({p1}),ISNUMBER({p2})),({w1}*{p1}+{w2}*{p2})/({w1}+{w2}),IF(ISNUMBER({p1}),{p1},{p2}))),'
                        f'{S["psmanual"]})'))
    setcell("check", (f'=IF({S["m1"]}="","先月データ未貼付",'
                      f'IF({S["m2"]}="","先々月データ未貼付（先月のみで予測中）",'
                      f'IF(EDATE({S["m2"]},1)<>{S["m1"]},"⚠ 先々月が先月の前月になっていません",'
                      f'IF({S["mp"]}="","OK（前年当月データなし：季節係数は手入力値）",'
                      f'IF(EDATE({S["mp"]},12)<>{S["cur"]},"⚠ 前年当月データが当月の1年前になっていません","OK")))))'))

    # ---------------- 集計：時刻構成比 ----------------
    HR = POOL + 5  # 時刻見出し行
    nH = len(HOURS)
    lastH = CL(1 + nH)
    c.cell(row=HR - 1, column=1, value="日タイプ×時刻の1日あたり平均客数（ランチ時間帯はランチ営業日のみ、ディナー時間帯は営業日の平均）  ※時刻は見出しを書き換え可").font = FB
    style(c.cell(row=HR, column=1, value="タイプ"), font=FW, fill=FILL_HD, align=CENTER)
    for j, h in enumerate(HOURS):
        style(c.cell(row=HR, column=2 + j, value=h), font=FW, fill=FILL_HD, fmt='0"時"', align=CENTER)
    for j, lab in enumerate(["ランチ計", "ディナー計"]):
        style(c.cell(row=HR, column=2 + nH + j, value=lab), font=FW, fill=FILL_HD, align=CENTER)
    AV = HR + 1  # 平均行（8タイプ + 全タイプ合算）
    lsum_col, dsum_col = CL(2 + nH), CL(3 + nH)
    for i, t in enumerate(TYPES + ["全体"]):
        r = AV + i
        trow = TR + i
        c.cell(row=r, column=1, value=t)
        style(c.cell(row=r, column=1), align=CENTER, fill=FILL_GR if t == "全体" else None)
        for j in range(nH):
            hl = CL(2 + j)
            hc = f"{hl}${HR}"
            per = []
            for blk, sheet, mkey, ncol, nlcol in ((B1, DS1, "m1", "B", "C"), (B2, DS2, "m2", "F", "G")):
                mc = S[mkey]
                base = (f"'{sheet}'!$Q:$Q,'{sheet}'!$K:$K,{hc},"
                        f"'{sheet}'!$A:$A,\">=\"&{mc},'{sheet}'!$A:$A,\"<=\"&EOMONTH({mc},0)")
                tcrit = "" if t == "全体" else f",'{sheet}'!$Z:$Z,$A{r}"
                if t == "全体":
                    n_d = f'COUNTIF({rng(bcol(blk, "客数"))},">0")'
                    n_l = f'COUNTIF({rng(bcol(blk, "ランチ営業"))},1)'
                else:
                    n_d, n_l = f"${ncol}${trow}", f"${nlcol}${trow}"
                n = f"IF({hc}<{dn},{n_l},{n_d})"
                val = (f"IF({hc}<{dn},SUMIFS({base}{tcrit},'{sheet}'!$AA:$AA,1),SUMIFS({base}{tcrit}))/{n}")
                per.append((f"IF(OR({mc}=\"\",{n}=0),\"\",{val})", n))
            (v1, n1), (v2, n2) = per
            c.cell(row=r, column=2 + j, value="=" + wavg(v1, f"N({n1})", v2, f"N({n2})"))
            style(c.cell(row=r, column=2 + j), font=FG, fmt="0.0")
        c.cell(row=r, column=2 + nH, value=f'=SUMIF($B${HR}:${lastH}${HR},"<"&{dn},B{r}:{lastH}{r})')
        c.cell(row=r, column=3 + nH, value=f'=SUMIF($B${HR}:${lastH}${HR},">="&{dn},B{r}:{lastH}{r})')
        for col in (2 + nH, 3 + nH):
            style(c.cell(row=r, column=col), font=FG, fmt="0.0")
    ALL = AV + 8
    sun_av = AV + TYPES.index("日")
    SH = ALL + 3
    c.cell(row=SH - 1, column=1, value="日タイプ×時刻の構成比（ランチ時間帯の合計＝100%、ディナー時間帯の合計＝100%）").font = FB
    style(c.cell(row=SH - 1, column=1, value="タイプ"), font=FW, fill=FILL_HD, align=CENTER)
    for j in range(nH):
        style(c.cell(row=SH - 1, column=2 + j, value=f"={CL(2 + j)}{HR}"), font=FW, fill=FILL_HD, fmt='0"時"', align=CENTER)
    c.cell(row=SH - 2, column=1, value="日タイプ×時刻の構成比（ランチ時間帯の合計＝100%、ディナー時間帯の合計＝100%）").font = FB
    for i, t in enumerate(TYPES):
        r = SH + i
        ra = AV + i
        c.cell(row=r, column=1, value=t)
        style(c.cell(row=r, column=1), align=CENTER)
        for j in range(nH):
            hl = CL(2 + j)
            hc = f"{hl}${HR}"
            lunch = (f'IF(N(${lsum_col}{ra})>0,N({hl}{ra})/${lsum_col}{ra},'
                     f'IF(N(${lsum_col}${ALL})>0,N({hl}${ALL})/${lsum_col}${ALL},0))')
            dinner = (f'IF(N(${dsum_col}{ra})>0,N({hl}{ra})/${dsum_col}{ra},'
                      f'IF(AND($A{r}="祝",N(${dsum_col}${sun_av})>0),N({hl}${sun_av})/${dsum_col}${sun_av},'
                      f'IF(N(${dsum_col}${ALL})>0,N({hl}${ALL})/${dsum_col}${ALL},0)))')
            c.cell(row=r, column=2 + j, value=f"=IF({hc}<{dn},{lunch},{dinner})")
            style(c.cell(row=r, column=2 + j), fmt="0.0%")
    # ---------------- 集計：日タイプ×時間帯の客単価 ----------------
    PT = SH + 11  # 時間帯名の見出し行（その下にサブ見出し、データ行）
    c.cell(row=PT - 1, column=1, value=f"日タイプ×時間帯の客単価（売上÷客数。客数が最低客数未満の月は使わず、全体→日タイプ全体の単価で代用）").font = FB
    style(c.cell(row=PT, column=1, value="タイプ"), font=FW, fill=FILL_HD, align=CENTER)
    style(c.cell(row=PT + 1, column=1), font=FW, fill=FILL_HD)
    for k, band in enumerate(BANDS):
        c0 = 2 + 3 * k
        c.merge_cells(start_row=PT, start_column=c0, end_row=PT, end_column=c0 + 2)
        style(c.cell(row=PT, column=c0, value=band), font=FW, fill=FILL_HD, align=CENTER)
        for j, lab in enumerate(["先月", "先々月", "採用"]):
            style(c.cell(row=PT + 1, column=c0 + j, value=lab), font=FW, fill=FILL_HD, align=CENTER)
    P0 = PT + 2  # データ先頭行（8タイプ）、P0+8 が全体
    pmin = S["pmin"]

    def band_price(sheet, mkey, band, tcrit):
        mc = S[mkey]
        crit = (f"'{sheet}'!$AB:$AB,\"{band}\",'{sheet}'!$A:$A,\">=\"&{mc},"
                f"'{sheet}'!$A:$A,\"<=\"&EOMONTH({mc},0){tcrit}")
        den = f"SUMIFS('{sheet}'!$Q:$Q,{crit})"
        return f'=IF({mc}="","",IF({den}<{pmin},"",SUMIFS(\'{sheet}\'!$L:$L,{crit})/{den}))'

    sun_p = P0 + TYPES.index("日")
    for i, t in enumerate(TYPES + ["全体"]):
        r = P0 + i
        c.cell(row=r, column=1, value=t)
        style(c.cell(row=r, column=1), align=CENTER, fill=FILL_GR if t == "全体" else None)
        for k, band in enumerate(BANDS):
            c0 = 2 + 3 * k
            x1, x2, ad = CL(c0), CL(c0 + 1), CL(c0 + 2)
            for j, (sheet, mkey) in enumerate(((DS1, "m1"), (DS2, "m2"))):
                tcrit = "" if t == "全体" else f",'{sheet}'!$Z:$Z,$A{r}"
                c.cell(row=r, column=c0 + j, value=band_price(sheet, mkey, band, tcrit))
            if t == "全体":
                fb = '""'
            elif t == "祝":
                fb = f'IF({ad}${sun_p}<>"",{ad}${sun_p},IF({ad}${P0 + 8}<>"",{ad}${P0 + 8},$P${TR + i}))'
            else:
                fb = f'IF({ad}${P0 + 8}<>"",{ad}${P0 + 8},$P${TR + i})'
            c.cell(row=r, column=c0 + 2, value=(
                f'=IF(AND({x1}{r}<>"",{x2}{r}<>""),({w1}*{x1}{r}+{w2}*{x2}{r})/({w1}+{w2}),'
                f'IF({x1}{r}<>"",{x1}{r},IF({x2}{r}<>"",{x2}{r},{fb})))'))
            for j in range(3):
                style(c.cell(row=r, column=c0 + j), fmt="#,##0", font=FB if j == 2 else F,
                      fill=FILL_GR if t == "全体" else None)

    c.column_dimensions["A"].width = 13
    for col in range(2, 32):
        c.column_dimensions[CL(col)].width = 10
    c.freeze_panes = "B3"

    # ---------------- 予測 ----------------
    f = ws_fc
    f.sheet_properties.tabColor = "70AD47"
    f["A1"] = f'=IF({S["cur"]}="","日別客数予測（先月データを貼り付けてください）",TEXT({S["cur"]},"yyyy年m月")&"　日別客数予測")'
    f["A1"].font = FT
    f["A2"] = (f'="データ: "&{S["check"]}&"　／　季節係数 "&TEXT({S["season"]},"0.000")&"　直近前年比 "&TEXT({S["yoy"]},"0.000")'
               f'&"　トレンド係数 "&TEXT({S["trend"]},"0.00")&"　前年同曜日の重み "&TEXT({S["wp"]},"0.00")')
    f["A2"].font = FNOTE
    FR0, FR1 = 9, 39
    kpis = [("月間予測客数", f"=SUM(O{FR0}:O{FR1})", "#,##0"), ("月間予測売上", f"=SUM(P{FR0}:P{FR1})", "#,##0"),
            ("先月実績客数", f"=SUM({rng(bcol(B1, '客数'))})", "#,##0"), ("先月比(客数)", '=IFERROR(D3/D5,"")', "0.0%"),
            ("予測客単価", '=IFERROR(D4/D3,"")', "#,##0")]
    for k, (lab, fml, fmt) in enumerate(kpis):
        r = 3 + k
        for col in range(1, 6):
            style(f.cell(row=r, column=col), font=FB, fill=FILL_KPI)
        f.cell(row=r, column=1, value=lab)
        f.cell(row=r, column=4, value=fml).number_format = fmt
        f.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        f.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
    notes = ["黄色の列だけ入力します：",
             "・営業区分（変更時のみ）… 通常は空欄＝土日祝はランチ+ディナー、平日はディナーのみ。夏休みの平日ランチ営業・臨時休業などはここで指定",
             "・日別補正 … イベント・天候・販促などの倍率（例：雨予報0.9、地域祭り1.2）"]
    for k, t in enumerate(notes):
        f.cell(row=3 + k, column=7, value=t).font = FNOTE
    hdr = ["日付", "曜日", "日タイプ", "営業区分\n(変更時のみ)", "適用\n営業区分", "ランチ\n基準", "ディナー\n基準",
           "前年\n同曜日", "前年\n客数", "季節\n係数", "トレンド\n係数", "日別\n補正", "予測\nランチ", "予測\nディナー",
           "予測客数", "予測売上", "予測\n客単価", "メモ"]
    header(f, FR0 - 1, 1, hdr)
    f.row_dimensions[FR0 - 1].height = 30
    helper_hdr = ["(計算用)前年ランチ", "(計算用)前年ディナー", "(計算用)前年タイプ", "(計算用)前年を使う"]
    header(f, FR0 - 1, 20, helper_hdr, fill=FILL_GR, font=FG)
    pa, pq, pk = f"'{DSP}'!$A:$A", f"'{DSP}'!$Q:$Q", f"'{DSP}'!$K:$K"
    dtype = f"集計!$A${TR}:$A${TR + 7}"
    wp, yoy, thr = S["wp"], S["yoy"], S["thr"]
    for i in range(31):
        r = FR0 + i
        A = f"$A{r}"
        v = {
            1: f'=IF(OR({S["cur"]}="",{i + 1}>{S["days"]}),"",{S["cur"]}+{i})',
            2: f'=IF({A}="","",{wd_formula(A)})',
            3: f'=IF({A}="","",{type_formula(A)})',
            4: None,
            5: f'=IF({A}="","",IF(D{r}<>"",D{r},IF(OR(C{r}="土",C{r}="日",C{r}="祝"),"{LD}","{DO}")))',
            6: f'=IF({A}="","",IF(E{r}="{LD}",INDEX(集計!$M${TR}:$M${TR + 7},MATCH(C{r},{dtype},0)),0))',
            7: f'=IF({A}="","",IF(E{r}="{CL_}",0,INDEX(集計!$L${TR}:$L${TR + 7},MATCH(C{r},{dtype},0))))',
            8: f'=IF(W{r}=1,{A}-364,"")',
            9: f'=IF(W{r}=1,IF(AND(E{r}="{LD}",T{r}>={thr}),T{r},0)+U{r},"")',
            10: f'=IF({A}="","",{S["season"]})',
            11: f'=IF({A}="","",{S["trend"]})',
            12: 1,
            13: (f'=IF({A}="","",IF(E{r}<>"{LD}",0,IF(AND(W{r}=1,T{r}>={thr}),(1-{wp})*F{r}*J{r}+{wp}*T{r}*{yoy},F{r}*J{r})*K{r}*N(L{r})))'),
            14: (f'=IF({A}="","",IF(E{r}="{CL_}",0,IF(W{r}=1,(1-{wp})*G{r}*J{r}+{wp}*U{r}*{yoy},G{r}*J{r})*K{r}*N(L{r})))'),
            15: f'=IF({A}="","",ROUND(M{r}+N{r},0))',
            16: f"=IF({A}=\"\",\"\",'{DSO}'!S{4 + i})",
            17: f'=IF({A}="","",IFERROR(P{r}/O{r},""))',
            18: None,
            20: (f'=IF({A}="","",IF(COUNTIF({pa},{A}-364)=0,"",'
                 f'SUMIFS({pq},{pa},{A}-364,{pk},">=6",{pk},"<"&{dn})))'),
            21: f'=IF(T{r}="","",SUMIFS({pq},{pa},{A}-364)-T{r})',
            22: f'=IF(T{r}="","",{type_formula(f"({A}-364)")})',
            23: (f'=IF({A}="",0,IF(AND({wp}>0,N(U{r})>0,C{r}<>"祝",V{r}<>"祝",E{r}<>"{CL_}"),1,0))'),
        }
        fmts = {1: "m/d", 6: "#,##0.0", 7: "#,##0.0", 8: "m/d(aaa)", 9: "#,##0", 10: "0.000", 11: "0.00", 12: "0.00",
                13: "#,##0.0", 14: "#,##0.0", 15: "#,##0", 16: "#,##0", 17: "#,##0"}
        for col in list(range(1, 19)) + [20, 21, 22, 23]:
            val = v.get(col)
            cell = f.cell(row=r, column=col, value=val)
            is_in = col in (4, 12, 18)
            font = FB if col == 15 else (FIN if is_in else (FG if col >= 20 else F))
            style(cell, font=font, fill=FILL_IN if is_in else None, fmt=fmts.get(col),
                  align=CENTER if col in (1, 2, 3, 4, 5, 8) else None)
    r = FR1 + 1
    for col in range(1, 19):
        style(f.cell(row=r, column=col), font=FB, fill=FILL_SUB, fmt="#,##0")
    f.cell(row=r, column=1, value="合計")
    for col in (13, 14, 15, 16):
        L = CL(col)
        f.cell(row=r, column=col, value=f"=SUM({L}{FR0}:{L}{FR1})")
    f.cell(row=r, column=17, value=f'=IFERROR(P{r}/O{r},"")')
    f.conditional_formatting.add(f"A{FR0}:C{FR1}", FormulaRule(formula=[f'$C{FR0}="土"'], font=Font(color="0070C0", bold=True)))
    f.conditional_formatting.add(f"A{FR0}:C{FR1}", FormulaRule(formula=[f'OR($C{FR0}="日",$C{FR0}="祝")'], font=Font(color="C00000", bold=True)))
    f.conditional_formatting.add(f"E{FR0}:Q{FR1}", FormulaRule(formula=[f'$E{FR0}="{CL_}"'], fill=PatternFill("solid", fgColor="D9D9D9")))
    f.conditional_formatting.add(f"E{FR0}:E{FR1}", FormulaRule(formula=[f'AND($D{FR0}<>"",$E{FR0}<>"{CL_}")'], font=Font(color="C55A11", bold=True)))
    dv = DataValidation(type="list", formula1=f'"{LD},{DO},{CL_}"', allow_blank=True)
    f.add_data_validation(dv)
    dv.add(f"D{FR0}:D{FR1}")
    dv2 = DataValidation(type="decimal", operator="between", formula1="0", formula2="5", allow_blank=True,
                         error="0〜5の倍率で入力してください（例：1.2）")
    f.add_data_validation(dv2)
    dv2.add(f"L{FR0}:L{FR1}")
    f.cell(row=FR0 - 1, column=12).comment = Comment("1.00=補正なし。1.20=2割増、0.80=2割減。", "tool")
    f.cell(row=FR0 - 1, column=8).comment = Comment("前年当月データの364日前（同じ曜日）。祝日が絡む日・休業日は前年の値を使わず空欄になります。", "tool")
    widths = [8, 5, 7, 13, 13, 8, 8, 10, 7, 7, 8, 7, 8, 8, 9, 11, 8, 24, 2, 9, 9, 9, 9]
    for col, w in enumerate(widths, start=1):
        f.column_dimensions[CL(col)].width = w
    f.freeze_panes = f"D{FR0}"
    for col in ("T", "U", "V", "W"):
        f.column_dimensions[col].hidden = True
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.title = "日別 予測客数（ランチ／ディナー）"
    ch.height, ch.width = 9, 24
    ch.add_data(Reference(f, min_col=13, max_col=14, min_row=FR0 - 1, max_row=FR1), titles_from_data=True)
    ch.set_categories(Reference(f, min_col=1, min_row=FR0, max_row=FR1))
    ch.x_axis.number_format = "d"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = "b"
    f.add_chart(ch, f"A{FR1 + 3}")

    # ---------------- 時間帯別予測 ----------------
    h = ws_hr
    h.sheet_properties.tabColor = "70AD47"
    h["A1"] = f'=IF({S["cur"]}="","時間帯別 予測客数",TEXT({S["cur"]},"yyyy年m月")&"　時間帯別 予測客数（シフト作成用）")'
    h["A1"].font = FT
    h["A2"] = "予測ランチ・予測ディナーを、同じ日タイプの時刻別構成比で配分。ディナーのみの日は夜営業開始前が0になります。小数は四捨五入表示のため合計が±数名ずれることがあります。"
    h["A2"].font = FNOTE
    header(h, 4, 1, ["日付", "タイプ", "営業区分", "日予測"])
    h.cell(row=3, column=4, value="時間帯→").font = FG
    for j in range(nH):
        style(h.cell(row=4, column=5 + j, value=f"=集計!{CL(2 + j)}{HR}"), font=FW, fill=FILL_HD, fmt='0"時"', align=CENTER)
        style(h.cell(row=3, column=5 + j, value=f"={band_formula(CL(5 + j) + '4')}"), font=FG, fill=FILL_GR,
              align=Alignment(horizontal="center", shrink_to_fit=True))
    for i in range(31):
        r = 5 + i
        fr = FR0 + i
        h.cell(row=r, column=1, value=f"=予測!A{fr}")
        h.cell(row=r, column=2, value=f"=予測!C{fr}")
        h.cell(row=r, column=3, value=f"=予測!E{fr}")
        h.cell(row=r, column=4, value=f"=予測!O{fr}")
        style(h.cell(row=r, column=1), fmt="m/d(aaa)", align=CENTER)
        style(h.cell(row=r, column=2), align=CENTER)
        style(h.cell(row=r, column=3), font=Font(name=FONT, size=8), align=CENTER)
        style(h.cell(row=r, column=4), fmt="#,##0", font=FB)
        for j in range(nH):
            hl = CL(2 + j)
            share = f"INDEX(集計!{hl}${SH}:{hl}${SH + 7},MATCH($B{r},集計!$A${SH}:$A${SH + 7},0))"
            h.cell(row=r, column=5 + j, value=(
                f'=IF($A{r}="","",IF(集計!{hl}${HR}<{dn},予測!$M{fr},予測!$N{fr})*{share})'))
            style(h.cell(row=r, column=5 + j), fmt="0")
    last_h = CL(4 + nH)
    h.conditional_formatting.add(f"E5:{last_h}35", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                                  end_type="max", end_color="F4B183"))
    h.conditional_formatting.add("A5:B35", FormulaRule(formula=['$B5="土"'], font=Font(color="0070C0", bold=True)))
    h.conditional_formatting.add("A5:B35", FormulaRule(formula=['OR($B5="日",$B5="祝")'], font=Font(color="C00000", bold=True)))
    h.column_dimensions["A"].width = 11
    h.column_dimensions["B"].width = 6
    h.column_dimensions["C"].width = 12
    h.column_dimensions["D"].width = 8
    for j in range(nH):
        h.column_dimensions[CL(5 + j)].width = 6
    h.freeze_panes = "E5"

    # ---------------- 売上予測貼付用（【事前】売上予測 と同じ行・列配置） ----------------
    o = ws_out
    o.sheet_properties.tabColor = "C00000"
    o["A1"] = f'=IF({S["cur"]}="","売上予測 貼り付け用",TEXT({S["cur"]},"yyyy年m月")&"　売上予測 貼り付け用")'
    o["A1"].font = FT
    o["V2"] = "【使い方】"
    o["V3"] = "① このシートの F4:M34（客単価4列＋客数4列）をコピー"
    o["V4"] = "② 【事前】売上予測 タブの F4 を選択 → 右クリック →「値の貼り付け」"
    o["V5"] = "③ 前年欄も使う場合は A4:C34 を同様に A4 へ（前年当月データ貼付時のみ）"
    o["V7"] = "※ 行・列の位置は【事前】売上予測と同じです（4行目＝1日）"
    o["V8"] = "※ 客単価は10円単位に丸め。客数は時間帯ごとに四捨五入し、端数差はディナーで調整（合計＝予測客数）"
    o["V9"] = "※ 時刻→時間帯の対応は「設定」シートの表で変更できます"
    o["V2"].font = FB
    for k in range(3, 10):
        o[f"V{k}"].font = FNOTE
    # 見出し（元シートと同じ結合）
    for rng_, val in (("A2:A3", "日"), ("B2:C2", "前年"), ("D2:D3", "日"), ("E2:E3", "曜"), ("F2:I2", "客単価"),
                      ("J2:M2", "時間帯別客数"), ("N2:Q2", "時間帯別売上"), ("R2:R3", "客数"), ("S2:S3", "売上"), ("T2:T3", "客単価")):
        o.merge_cells(rng_)
        o[rng_.split(":")[0]] = val
    o["B3"], o["C3"] = "客数", "売上"
    for k, band in enumerate(BANDS):
        for c0 in (6, 10, 14):
            o.cell(row=3, column=c0 + k, value=band)
    for row in o.iter_rows(min_row=2, max_row=3, min_col=1, max_col=20):
        for cc in row:
            style(cc, font=FW, fill=FILL_HD, align=CENTER)
    for col in range(6, 14):
        for rr in (2, 3):
            o.cell(row=rr, column=col).fill = PatternFill("solid", fgColor="C00000")
    band_hdr = f"時間帯別予測!$E$3:${CL(4 + nH)}$3"
    for i in range(31):
        r = 4 + i
        fr = FR0 + i
        hr_row = 5 + i
        d = f"予測!$A${fr}"
        tp = f"予測!$C${fr}"
        v = {
            1: f'=IF({d}="","",IF(COUNTIF({pa},{d}-364)>0,DAY({d}-364),""))',
            2: f'=IF(A{r}="","",SUMIFS({pq},{pa},{d}-364))',
            3: f"=IF(A{r}=\"\",\"\",SUMIFS('{DSP}'!$L:$L,{pa},{d}-364))",
            4: f'=IF({d}="","",DAY({d}))',
            5: f'=IF({d}="","",{wd_formula(d)})',
        }
        for k, band in enumerate(BANDS):
            ad = CL(2 + 3 * k + 2)
            v[6 + k] = (f'=IF({d}="",0,ROUND(INDEX(集計!${ad}${P0}:${ad}${P0 + 7},MATCH({tp},集計!$A${P0}:$A${P0 + 7},0))'
                        f'*{S["pseason"]}*{S["ptrend"]},-1))')
            raw = f"SUMIF({band_hdr},\"{band}\",時間帯別予測!$E${hr_row}:${CL(4 + nH)}${hr_row})"
            if band == "ディナー":
                v[10 + k] = f'=IF({d}="",0,MAX(0,予測!$O${fr}-J{r}-K{r}-M{r}))'
            else:
                v[10 + k] = f'=IF({d}="",0,ROUND({raw},0))'
            v[14 + k] = f"={CL(6 + k)}{r}*{CL(10 + k)}{r}"
        v[18] = f"=SUM(J{r}:M{r})"
        v[19] = f"=SUM(N{r}:Q{r})"
        v[20] = f'=IFERROR(S{r}/R{r},"")'
        for col in range(1, 21):
            cell = o.cell(row=r, column=col, value=v[col])
            style(cell, fmt="#,##0", align=CENTER if col in (1, 4, 5) else None,
                  font=FB if col in (18, 19) else F,
                  fill=PatternFill("solid", fgColor="FCE4D6") if 6 <= col <= 13 else None)
    r = 35
    o.cell(row=r, column=1, value="合計")
    for col in list(range(2, 4)) + list(range(10, 20)):
        L = CL(col)
        o.cell(row=r, column=col, value=f"=SUM({L}4:{L}34)")
    for k in range(4):  # 時間帯別の平均客単価
        o.cell(row=r, column=6 + k, value=f'=IFERROR({CL(14 + k)}35/{CL(10 + k)}35,"")')
    o.cell(row=r, column=20, value='=IFERROR(S35/R35,"")')
    for col in range(1, 21):
        style(o.cell(row=r, column=col), font=FB, fill=FILL_SUB, fmt="#,##0")
    o.conditional_formatting.add("D4:E34", FormulaRule(formula=['$E4="土"'], font=Font(color="0070C0", bold=True)))
    o.conditional_formatting.add("D4:E34", FormulaRule(formula=['$E4="日"'],
                                                       font=Font(color="C00000", bold=True)))
    for col, w in zip("ABCDEFGHIJKLMNOPQRST", (5, 7, 10, 5, 5, 7, 7, 7, 7, 7, 7, 7, 7, 11, 11, 12, 11, 7, 12, 8)):
        o.column_dimensions[col].width = w
    o.column_dimensions["V"].width = 70
    o.freeze_panes = "F4"

    # ---------------- 使い方 ----------------
    u = ws_help
    u.sheet_properties.tabColor = "4472C4"
    u.column_dimensions["A"].width = 3
    u.column_dimensions["B"].width = 120
    lines = [
        ("日別客数予測ツール", FT),
        ("", F),
        ("■ 毎月の使い方", FB),
        ("① 「先月データ」シートのA1を選択 → 先月の時間帯別売上実績表（見出し行ごと）を貼り付け", F),
        ("② 「先々月データ」シートのA1に先々月分を貼り付け", F),
        ("③ （推奨）「前年当月データ」シートのA1に、1年前の当月＋翌月1日（例：当月が2026年10月なら2025/10/1〜11/1）を貼り付け", F),
        ("   ※ 貼り付け前に、前回のデータを A〜X 列ごと削除してください（Z列以降の自動計算列は消さない）", F),
        ("④ 「予測」シートで当月の営業カレンダーを確認。夏休みの平日ランチ営業・臨時休業などは黄色の「営業区分」で指定", F),
        ("⑤ 必要なら「日別補正」にイベント・天候などの倍率を入力。「時間帯別予測」はシフト作成に使えます", F),
        ("⑥ 「売上予測貼付用」の F4:M34 をコピー →【事前】売上予測 タブの F4 に「値の貼り付け」（客単価・客数が時間帯別に入ります）", F),
        ("", F),
        ("■ 予測の考え方", FB),
        ("予測客数 ＝ 予測ランチ ＋ 予測ディナー（ランチとディナーを別々に予測して足す）", FB),
        ("・営業区分：土日祝＝ランチ+ディナー、平日＝ディナーのみ が標準。ディナーのみの日はランチ分を0にする", F),
        ("・ディナー基準：先月・先々月の日タイプ別（月〜日＋祝）のディナー平均を重み 0.6:0.4 で加重平均（客数0の休業日は除外）", F),
        ("・ランチ基準：同じく、実際にランチ営業した日だけの平均（夏休み・お盆の平日ランチが平常月の平日に混ざらない）", F),
        ("・季節係数：前年当月データがあれば『前年当月 ÷ 前年の先月・先々月』の曜日別ディナー平均の比で自動計算（祝日除く）", F),
        ("・前年同曜日：前年当月データの364日前（同じ曜日）の実績 × 直近の前年比 を、直近ベースと 0.5:0.5 でブレンド", F),
        ("   祝日・祝前後のずれを避けるため、当日または前年の該当日が祝日の日は前年の値を使わない", F),
        ("・トレンド係数：販促・値上げ・競合出店など全体の上げ下げを手入力　・日別補正：その日だけの要因", F),
        ("・時間帯別：予測ランチ／ディナーを、同じ日タイプの時刻別構成比で配分", F),
        ("・時間帯別の客数：時間帯別予測を「設定」の時刻→時間帯対応表（ランチ11〜14時／アイドル15〜17時／ディナー18〜21時／ナイト22時〜）で合計", F),
        ("", F),
        ("■ 客単価の予測（客数と同じ考え方）", FB),
        ("予測客単価（時間帯別）＝ 日タイプ×時間帯の基準客単価 × 客単価の季節係数 × 客単価トレンド係数", FB),
        ("・基準客単価：先月・先々月の『売上÷客数』を日タイプ（月〜日＋祝）×時間帯ごとに計算し、0.6:0.4 で加重平均", F),
        ("   客数が少ない時間帯（最低客数未満）は単価がぶれるので、全曜日合算の時間帯単価 → 日タイプ全体の単価で代用", F),
        ("・客単価の季節係数：前年当月の客単価 ÷ 前年の先月・先々月の客単価（データの「売上前年」「客数前年」から計算）", F),
        ("・客単価トレンド係数：値上げ・メニュー改定などを手入力", F),
        ("・予測売上：時間帯別の 予測客数 × 予測客単価 の合計（日別の客単価はそこから逆算）", F),
        ("", F),
        ("■ シート一覧", FB),
        ("予測／時間帯別予測：結果　売上予測貼付用：【事前】売上予測への貼り付け元　設定：営業時間・重み・係数（黄色が入力欄）　各データ：貼り付け先　祝日：祝日・特異日リスト　集計：内部計算", F),
        ("", F),
        ("■ 色の意味", FB),
        ("黄色背景・青字＝入力してよいセル　／　それ以外＝自動計算（触らない）", F),
    ]
    for i, (txt, font) in enumerate(lines, start=2):
        cc = u.cell(row=i, column=2, value=txt)
        cc.font = font
        cc.alignment = Alignment(wrap_text=True, vertical="top")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title in (DS1, DS2, DSP)
    wb.active = 1
    wb.save(out)


if __name__ == "__main__":
    args = sys.argv[2:5] + [None, None, None]
    build(sys.argv[1], *[a or None for a in args[:3]])
